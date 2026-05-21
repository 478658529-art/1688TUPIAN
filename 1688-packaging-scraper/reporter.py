"""
报告生成模块
生成 HTML 网页报告和 Excel 表格
"""
from datetime import datetime
from pathlib import Path
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import config

# image_analyzer might not be available if opencv not installed
try:
    from image_analyzer import summarize_analysis
    HAS_ANALYZER = True
except ImportError:
    HAS_ANALYZER = False


def _build_analysis_section(products: list) -> str:
    """构建图片视觉分析报告HTML片段"""
    if not HAS_ANALYZER:
        return ""
    try:
        summary = summarize_analysis(products)
    except:
        return ""

    if summary.get("total_analyzed", 0) == 0:
        return ""

    analyzed = summary["total_analyzed"]

    # ── 风格分布 ──
    cat_html = ""
    for cat, count in summary.get("visual_categories", {}).items():
        pct = round(count / analyzed * 100, 1)
        cat_html += f"""
        <div class="ana-stat-item">
            <div class="ana-stat-name">{cat}</div>
            <div class="ana-stat-bar"><div class="ana-stat-fill" style="width:{pct}%"></div></div>
            <div class="ana-stat-val">{count}件 ({pct}%)</div>
        </div>"""

    # ── 热门标签排行 ──
    tags_html = ""
    for tag, count in summary.get("style_tags_rank", {}).items():
        tags_html += f'<span class="ana-tag">{tag} <small>x{count}</small></span> '

    # ── 颜色分布 ──
    color_html = ""
    for color_name, pct in summary.get("color_distribution", {}).items():
        color_html += f'<span class="ana-tag color-tag">{color_name} <small>{pct}%</small></span> '

    # ── 材质分布 ──
    mat_html = ""
    for mat, count in summary.get("material_distribution", {}).items():
        mat_html += f'<span class="ana-tag mat-tag">{mat} <small>x{count}</small></span> '

    # ── 爆款特征排行 ──
    trend_html = ""
    for feat, count in summary.get("trending_features", {}).items():
        bar_width = min(round(count / analyzed * 100, 1) * 2, 100)
        trend_html += f"""
        <div class="ana-stat-item">
            <div class="ana-stat-name">{feat}</div>
            <div class="ana-stat-bar"><div class="ana-stat-fill trend-fill" style="width:{bar_width}%"></div></div>
            <div class="ana-stat-val">{count}/{analyzed}</div>
        </div>"""

    # ── 形状分布 ──
    shape_html = ""
    for shp, count in summary.get("shape_distribution", {}).items():
        shape_html += f'<span class="ana-tag shape-tag">{shp} <small>x{count}</small></span> '

    return f"""
    <div class="analysis-report">
        <div class="ana-header">
            <h2>🎨 图片视觉分析报告</h2>
            <span class="ana-subtitle">AI 自动分析 {analyzed} 张产品图片 · 综合归类总结</span>
        </div>

        <div class="ana-grid">
            <!-- 视觉风格分布 -->
            <div class="ana-card">
                <div class="ana-card-title">📊 视觉风格分布</div>
                <div class="ana-stat-list">{cat_html}</div>
            </div>

            <!-- 爆款特征趋势 -->
            <div class="ana-card">
                <div class="ana-card-title">🔥 爆款特征趋势</div>
                <div class="ana-stat-list">{trend_html}</div>
            </div>

            <!-- 风格标签 -->
            <div class="ana-card">
                <div class="ana-card-title">🏷️ 设计风格标签</div>
                <div class="ana-tag-wrap">{tags_html}</div>
            </div>

            <!-- 颜色分布 -->
            <div class="ana-card">
                <div class="ana-card-title">🎨 主色调分布</div>
                <div class="ana-tag-wrap">{color_html}</div>
            </div>

            <!-- 材质分布 -->
            <div class="ana-card">
                <div class="ana-card-title">🧵 材质分析</div>
                <div class="ana-tag-wrap">{mat_html}</div>
            </div>

            <!-- 形状分布 -->
            <div class="ana-card">
                <div class="ana-card-title">📐 形状分析</div>
                <div class="ana-tag-wrap">{shape_html}</div>
            </div>
        </div>
    </div>
    """


def generate_analysis_report(products: list, output_path: Path):
    """生成独立的图片视觉分析报告（JSON格式）"""
    if not HAS_ANALYZER:
        print("⚠️  图片分析器未安装，跳过分析报告")
        return None
    try:
        summary = summarize_analysis(products)
        import json
        summary_path = output_path.with_suffix(".json")
        summary_path.write_text(
            json.dumps(summary, ensure_ascii=False, indent=2),
            encoding="utf-8"
        )
        print(f"📊 图片分析统计已保存：{summary_path}")
        return summary
    except Exception as e:
        print(f"⚠️  生成分析报告出错：{e}")
        return None


def generate_html_report(products: list, output_path: Path):
    """生成图文并茂的HTML报告（含图片视觉分析）"""
    today = datetime.now().strftime("%Y-%m-%d")
    weekday_map = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]
    weekday = weekday_map[datetime.now().weekday()]

    # 按分类汇总
    categorized = defaultdict(list)
    for p in products:
        categorized[p.get("category", "其他包装")].append(p)

    total = len(products)
    categories_count = len(categorized)

    category_cards = ""
    for cat, items in sorted(categorized.items(), key=lambda x: -len(x[1])):
        count = len(items)
        items_html = ""
        for p in items:
            title = p.get("title", "")
            price = p.get("price", "面议")
            shop = p.get("shop_name", "")
            keyword = p.get("source_keyword", "")
            sales = p.get("sales_info", "")
            detail = p.get("detail_text", "")[:200]
            img_local = p.get("main_image_local", "")
            img_url = p.get("main_image_url", "")
            product_url = p.get("url", "")
            specs = p.get("specs", {})

            # 图片
            if img_local and Path(img_local).exists():
                rel_path = Path(img_local).relative_to(config.BASE_DIR)
                img_tag = f'<img src="../{rel_path.as_posix()}" alt="{title}" class="product-img">'
            elif img_url:
                img_tag = f'<img src="{img_url}" alt="{title}" class="product-img">'
            else:
                img_tag = '<div class="no-img">暂无图片</div>'

            specs_html = ""
            if specs:
                specs_html = '<div class="specs"><small>'
                for k, v in list(specs.items())[:5]:
                    specs_html += f"<span>{k}: {v}</span> "
                specs_html += "</small></div>"

            items_html += f"""
            <div class="product-card">
                <a href="{product_url}" target="_blank" class="img-link">
                    {img_tag}
                </a>
                <div class="product-info">
                    <h3 class="product-title">
                        <a href="{product_url}" target="_blank">{title[:50]}{'...' if len(title) > 50 else ''}</a>
                    </h3>
                    <div class="product-price">💰 {price}</div>
                    <div class="product-shop">🏪 {shop}</div>
                    <div class="product-keyword">🔑 来源词：{keyword}</div>
                    {f'<div class="product-sales">📊 {sales}</div>' if sales else ''}
                    {specs_html}
                    <div class="product-detail"><small>{detail}</small></div>
                </div>
            </div>
            """

        category_cards += f"""
        <div class="category-section">
            <div class="category-header">
                <h2>{cat}</h2>
                <span class="category-count">{count} 件产品</span>
            </div>
            <div class="product-grid">
                {items_html}
            </div>
        </div>
        """

    # ── 构建分析报告 HTML（如果启用） ──
    if config.GENERATE_ANALYSIS_REPORT and HAS_ANALYZER:
        analysis_section = _build_analysis_section(products)
    else:
        analysis_section = ""

    html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{config.REPORT_TITLE} - {today}</title>
    <style>
        * {{ margin: 0; padding: 0; box-sizing: border-box; }}
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
            background: #f5f5f5; color: #333; line-height: 1.6;
        }}
        .header {{
            background: linear-gradient(135deg, #FF6A00, #FF4500);
            color: white; padding: 30px 20px; text-align: center;
        }}
        .header h1 {{ font-size: 28px; margin-bottom: 5px; }}
        .header .subtitle {{ font-size: 14px; opacity: 0.9; }}
        .stats-bar {{
            display: flex; justify-content: center; gap: 30px;
            padding: 15px; background: white; border-bottom: 1px solid #eee;
        }}
        .stat-item {{ text-align: center; }}
        .stat-number {{ font-size: 24px; font-weight: bold; color: #FF6A00; }}
        .stat-label {{ font-size: 12px; color: #888; }}
        .container {{ max-width: 1400px; margin: 0 auto; padding: 20px; }}
        .category-section {{ margin-bottom: 30px; }}
        .category-header {{
            display: flex; justify-content: space-between; align-items: center;
            padding: 12px 20px; background: white; border-radius: 10px;
            margin-bottom: 15px; box-shadow: 0 2px 8px rgba(0,0,0,0.06);
        }}
        .category-header h2 {{ font-size: 18px; color: #333; }}
        .category-count {{ font-size: 13px; color: #FF6A00; font-weight: bold; }}
        .product-grid {{ display: grid; grid-template-columns: repeat(auto-fill, minmax(320px, 1fr)); gap: 15px; }}
        .product-card {{ background: white; border-radius: 10px; overflow: hidden; box-shadow: 0 2px 8px rgba(0,0,0,0.06); transition: transform 0.2s; }}
        .product-card:hover {{ transform: translateY(-3px); box-shadow: 0 4px 16px rgba(0,0,0,0.12); }}
        .img-link {{ display: block; width: 100%; height: 220px; overflow: hidden; }}
        .product-img {{ width: 100%; height: 100%; object-fit: cover; transition: transform 0.3s; }}
        .product-card:hover .product-img {{ transform: scale(1.05); }}
        .no-img {{ width: 100%; height: 220px; background: #f0f0f0; display: flex; align-items: center; justify-content: center; color: #ccc; }}
        .product-info {{ padding: 12px 15px 15px; }}
        .product-title {{ font-size: 14px; margin-bottom: 8px; line-height: 1.4; }}
        .product-title a {{ color: #333; text-decoration: none; }}
        .product-title a:hover {{ color: #FF6A00; }}
        .product-price {{ font-size: 16px; font-weight: bold; color: #FF4500; margin-bottom: 4px; }}
        .product-shop, .product-keyword, .product-sales {{ font-size: 12px; color: #888; margin-bottom: 2px; }}
        .specs {{ margin: 6px 0; }}
        .specs span {{ display: inline-block; background: #f0f6ff; color: #0077CC; padding: 2px 8px; border-radius: 4px; font-size: 11px; margin: 2px; }}
        .product-detail {{ margin-top: 8px; color: #666; max-height: 60px; overflow: hidden; }}
        .footer {{ text-align: center; padding: 20px; color: #aaa; font-size: 12px; }}
        /* 🎨 图片视觉分析报告样式 */
        .analysis-report {{ background: linear-gradient(135deg, #667eea 0%, #764ba2 100%); color: white; padding: 30px; }}
        .ana-header {{ text-align: center; margin-bottom: 25px; }}
        .ana-header h2 {{ font-size: 24px; margin-bottom: 5px; }}
        .ana-subtitle {{ font-size: 13px; opacity: 0.85; }}
        .ana-grid {{ display: grid; grid-template-columns: repeat(auto-fill, minmax(380px, 1fr)); gap: 18px; max-width: 1400px; margin: 0 auto; }}
        .ana-card {{ background: rgba(255,255,255,0.12); border-radius: 12px; padding: 18px; backdrop-filter: blur(10px); }}
        .ana-card-title {{ font-size: 15px; font-weight: bold; margin-bottom: 12px; opacity: 0.95; }}
        .ana-stat-item {{ display: flex; align-items: center; gap: 10px; margin-bottom: 8px; font-size: 13px; }}
        .ana-stat-name {{ flex-shrink: 0; min-width: 100px; }}
        .ana-stat-bar {{ flex-grow: 1; height: 18px; background: rgba(255,255,255,0.2); border-radius: 10px; overflow: hidden; }}
        .ana-stat-fill {{ height: 100%; background: linear-gradient(90deg, #f093fb, #f5576c); border-radius: 10px; transition: width 0.5s; }}
        .trend-fill {{ background: linear-gradient(90deg, #4facfe, #00f2fe); }}
        .ana-stat-val {{ flex-shrink: 0; min-width: 70px; text-align: right; font-size: 12px; opacity: 0.85; }}
        .ana-tag-wrap {{ display: flex; flex-wrap: wrap; gap: 6px; }}
        .ana-tag {{ display: inline-block; background: rgba(255,255,255,0.18); padding: 4px 10px; border-radius: 15px; font-size: 12px; }}
        .ana-tag small {{ opacity: 0.7; margin-left: 3px; }}
        .color-tag {{ border-left: 3px solid #f093fb; }}
        .mat-tag {{ border-left: 3px solid #4facfe; }}
        .shape-tag {{ border-left: 3px solid #43e97b; }}
        @media (max-width: 700px) {{ .product-grid {{ grid-template-columns: 1fr; }} .stats-bar {{ flex-wrap: wrap; }} .ana-grid {{ grid-template-columns: 1fr; }} }}
    </style>
</head>
<body>
    <div class="header">
        <h1>{config.REPORT_TITLE}</h1>
        <div class="subtitle">{today}（{weekday}） | 共 {total} 件包装类产品</div>
    </div>
    <div class="stats-bar">
        <div class="stat-item"><div class="stat-number">{total}</div><div class="stat-label">总产品数</div></div>
        <div class="stat-item"><div class="stat-number">{categories_count}</div><div class="stat-label">包装分类</div></div>
        <div class="stat-item"><div class="stat-number">{len(config.PACKAGING_KEYWORDS)}</div><div class="stat-label">搜索关键词</div></div>
    </div>
    {analysis_section}
    <div class="container">{category_cards}</div>
    <div class="footer">
        <p>数据来源：1688.com | 采集时间：{datetime.now().strftime("%Y-%m-%d %H:%M")}</p>
        <p>由 1688 包装采集工具 自动生成</p>
    </div>
</body>
</html>"""

    output_path.write_text(html, encoding="utf-8")
    print(f"📄 HTML 报告已生成：{output_path}")
    return output_path


def generate_excel_report(products: list, output_path: Path):
    """生成 Excel 表格报告"""
    wb = Workbook()
    ws = wb.active
    ws.title = "产品总览"
    today = datetime.now().strftime("%Y-%m-%d")

    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="FF6A00", end_color="FF6A00", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_font = Font(name="微软雅黑", size=10)
    cell_align = Alignment(vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )

    ws.merge_cells("A1:J1")
    ws["A1"].value = f"{config.REPORT_TITLE} - {today}"
    ws["A1"].font = Font(name="微软雅黑", size=16, bold=True, color="FF6A00")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    headers = ["序号", "分类", "产品标题", "价格", "店铺名称", "来源关键词", "销量信息", "主图(本地路径)", "详情摘要", "产品链接"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    for idx, p in enumerate(products, 1):
        row = idx + 3
        data = [
            idx,
            p.get("category", ""),
            p.get("title", ""),
            p.get("price", ""),
            p.get("shop_name", ""),
            p.get("source_keyword", ""),
            p.get("sales_info", ""),
            p.get("main_image_local", ""),
            (p.get("detail_text", "") or "")[:200],
            p.get("url", ""),
        ]
        for col_idx, value in enumerate(data, 1):
            cell = ws.cell(row=row, column=col_idx, value=value)
            cell.font = cell_font
            cell.alignment = cell_align
            cell.border = thin_border
        ws.row_dimensions[row].height = 40

    col_widths = [6, 14, 40, 12, 18, 14, 14, 35, 45, 50]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"

    # Sheet 2
    ws2 = wb.create_sheet("分类汇总")
    categorized = defaultdict(list)
    for p in products:
        categorized[p.get("category", "其他包装")].append(p)

    ws2.merge_cells("A1:D1")
    ws2["A1"].value = f"分类汇总 - {today}"
    ws2["A1"].font = Font(name="微软雅黑", size=14, bold=True, color="FF6A00")
    ws2["A1"].alignment = Alignment(horizontal="center")

    r = 3
    cat_headers = ["分类", "产品数量", "价格区间", "代表产品"]
    for c_idx, h in enumerate(cat_headers, 1):
        cell = ws2.cell(row=r, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    cat_fill_colors = ["FFF0F0", "FFF8F0", "F0F6FF", "F0FFF8", "F8F0FF", "FFF0F6", "F0F0F0", "FFFBE6"]

    for idx, (cat, items) in enumerate(sorted(categorized.items(), key=lambda x: -len(x[1]))):
        r += 1
        prices = [p.get("price", "") for p in items if p.get("price")]
        price_range = f"{min(prices)} ~ {max(prices)}" if prices else "面议"
        samples = "\n".join([p.get("title", "")[:30] for p in items[:3]])
        data = [cat, len(items), price_range, samples]
        fill = PatternFill(start_color=cat_fill_colors[idx % len(cat_fill_colors)], end_color=cat_fill_colors[idx % len(cat_fill_colors)], fill_type="solid")
        for c_idx, value in enumerate(data, 1):
            cell = ws2.cell(row=r, column=c_idx, value=value)
            cell.font = cell_font
            cell.alignment = cell_align
            cell.border = thin_border
            cell.fill = fill
        ws2.row_dimensions[r].height = 50

    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 12
    ws2.column_dimensions["C"].width = 22
    ws2.column_dimensions["D"].width = 60
    ws2.freeze_panes = "A4"

    wb.save(str(output_path))
    print(f"📊 Excel 报告已生成：{output_path}")
    return output_path
